"""
MallLite 自动化测试 - 数据读取模块
支持 JSON / YAML / CSV / Excel 四种格式

用法：
    from common.data_reader import DataReader

    # 读取 JSON
    users = DataReader.read_json("login_data.json")

    # 读取 YAML
    cases = DataReader.read_yaml("search_data.yaml")

    # 读取 CSV
    rows = DataReader.read_csv("cart_data.csv")

    # 读取 Excel
    data = DataReader.read_excel("user_data.xlsx", sheet_name="登录测试")
"""

import json
import csv
import yaml
from pathlib import Path
from common.logger import get_logger

logger = get_logger("data_reader")

# 测试数据目录
PROJECT_ROOT = Path(__file__).parent.parent
DEFAULT_DATA_DIR = PROJECT_ROOT / "test_data"


class DataReader:
    """数据读取工具类（所有方法都是静态方法，不需要实例化）"""

    @staticmethod
    def read_json(filename, data_dir=None):
        """
        读取 JSON 文件

        参数：
            filename: 文件名（如 "login_data.json"）或完整路径
            data_dir: 数据目录（默认 test_data/）

        返回：
            解析后的 Python 对象（dict 或 list）
        """
        filepath = DataReader._resolve_path(filename, data_dir)

        try:
            with open(filepath, "r", encoding="utf-8") as f:
                data = json.load(f)
            logger.debug(f"读取 JSON 成功：{filepath}（{len(data) if isinstance(data, list) else 'dict'}）")
            return data
        except FileNotFoundError:
            logger.error(f"JSON 文件不存在：{filepath}")
            raise
        except json.JSONDecodeError as e:
            logger.error(f"JSON 格式错误：{filepath}，{e}")
            raise

    @staticmethod
    def read_yaml(filename, data_dir=None):
        """
        读取 YAML 文件

        参数：
            filename: 文件名或完整路径
            data_dir: 数据目录

        返回：
            解析后的 Python 对象
        """
        filepath = DataReader._resolve_path(filename, data_dir)

        try:
            with open(filepath, "r", encoding="utf-8") as f:
                data = yaml.safe_load(f)
            logger.debug(f"读取 YAML 成功：{filepath}")
            return data
        except FileNotFoundError:
            logger.error(f"YAML 文件不存在：{filepath}")
            raise
        except yaml.YAMLError as e:
            logger.error(f"YAML 格式错误：{filepath}，{e}")
            raise

    @staticmethod
    def read_csv(filename, data_dir=None, encoding="utf-8"):
        """
        读取 CSV 文件

        参数：
            filename: 文件名或完整路径
            data_dir: 数据目录
            encoding: 编码格式

        返回：
            字典列表，每个字典是一行数据，键为表头
        """
        filepath = DataReader._resolve_path(filename, data_dir)

        try:
            with open(filepath, "r", encoding=encoding) as f:
                reader = csv.DictReader(f)
                rows = list(reader)
            logger.debug(f"读取 CSV 成功：{filepath}（{len(rows)} 行）")
            return rows
        except FileNotFoundError:
            logger.error(f"CSV 文件不存在：{filepath}")
            raise

    @staticmethod
    def read_excel(filename, sheet_name=None, data_dir=None):
        """
        读取 Excel 文件

        参数：
            filename: 文件名或完整路径
            sheet_name: 工作表名称（默认读第一个）
            data_dir: 数据目录

        返回：
            字典列表，每个字典是一行数据，键为表头
        """
        filepath = DataReader._resolve_path(filename, data_dir)

        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.error("openpyxl 未安装，请执行：pip install openpyxl")
            raise

        try:
            wb = load_workbook(str(filepath), read_only=True)
            ws = wb[sheet_name] if sheet_name else wb.active

            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                logger.warning(f"Excel 文件为空：{filepath}")
                return []

            # 第一行作为表头
            headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]
            result = []
            for row in rows[1:]:
                row_dict = {}
                for i, value in enumerate(row):
                    if i < len(headers):
                        row_dict[headers[i]] = value
                result.append(row_dict)

            wb.close()
            logger.debug(f"读取 Excel 成功：{filepath}（{len(result)} 行）")
            return result
        except FileNotFoundError:
            logger.error(f"Excel 文件不存在：{filepath}")
            raise
        except Exception as e:
            logger.error(f"读取 Excel 失败：{filepath}，{e}")
            raise

    @staticmethod
    def _resolve_path(filename, data_dir=None):
        """
        解析文件路径

        如果 filename 是绝对路径，直接使用
        如果是相对路径，在 data_dir 中查找（默认 test_data/）
        """
        filepath = Path(filename)

        # 如果是绝对路径且存在，直接返回
        if filepath.is_absolute() and filepath.exists():
            return filepath

        # 在数据目录中查找
        base_dir = Path(data_dir) if data_dir else DEFAULT_DATA_DIR
        full_path = base_dir / filename

        if full_path.exists():
            return full_path

        # 如果 data_dir 也没找到，尝试当前工作目录
        if filepath.exists():
            return filepath

        # 都没找到，返回默认路径（让调用方处理 FileNotFoundError）
        logger.warning(f"文件未找到：{full_path}")
        return full_path