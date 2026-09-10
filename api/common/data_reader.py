"""
测试数据读取工具

从 test_data/ 目录读取 JSON / YAML 文件
"""

import json
from pathlib import Path
from common.logger import get_logger

logger = get_logger("data_reader")

DATA_DIR = Path(__file__).parent.parent / "test_data"


def read_json(filename):
    """读取 JSON 文件"""
    filepath = DATA_DIR / filename
    if not filepath.exists():
        raise FileNotFoundError(f"数据文件不存在：{filepath}")
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)
    logger.debug(f"读取数据：{filename}")
    return data


def read_yaml(filename):
    """读取 YAML 文件"""
    import yaml
    filepath = DATA_DIR / filename
    if not filepath.exists():
        raise FileNotFoundError(f"数据文件不存在：{filepath}")
    with open(filepath, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    logger.debug(f"读取数据：{filename}")
    return data

# ===== 以下是 YAML 支持（如果之前没加的话） =====

def read_yaml(filename):
    """
    读取 YAML 数据文件

    参数：
        filename: 文件名（在 test_data 目录下）

    返回：
        解析后的 Python 对象（字典或列表）
    """
    import yaml
    filepath = DATA_DIR / filename
    if not filepath.exists():
        raise FileNotFoundError(f"数据文件不存在：{filepath}")
    with open(filepath, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    logger.debug(f"读取 YAML 数据：{filename}")
    return data


def read_excel(filename, sheet_name=None):
    """
    读取 Excel 测试数据

    参数：
        filename: Excel 文件名（在 test_data 目录下）
        sheet_name: Sheet 名称（None 读第一个 Sheet）

    返回：
        列表，每个元素是一个字典（第一行是表头）
        [
            {"case_name": "管理员", "username": "admin", ...},
            {"case_name": "普通用户", "username": "testuser", ...},
        ]

    知识点：为什么返回字典列表而不是二维列表？
        字典列表可以直接用 case["username"] 取值，比 case[1] 更直观。
        而且调换 Excel 列的顺序不会影响代码。
    """
    from openpyxl import load_workbook

    filepath = DATA_DIR / filename
    if not filepath.exists():
        raise FileNotFoundError(f"Excel 文件不存在：{filepath}")

    wb = load_workbook(filepath, read_only=True)

    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    # 第一行是表头
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []

    headers = [str(h) if h else f"col_{i}" for i, h in enumerate(rows[0])]

    # 从第二行开始是数据
    result = []
    for row in rows[1:]:
        # 跳过全空的行
        if all(cell is None for cell in row):
            continue
        item = dict(zip(headers, row))
        result.append(item)

    wb.close()
    logger.debug(f"读取 Excel：{filename}（Sheet: {sheet_name or 'active'}），共 {len(result)} 条")
    return result